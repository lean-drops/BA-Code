# /Users/programming/PycharmProjects/Find_Bibliography_NEw/insert_footnotes.py
"""
Nimmt ein markiertes DOCX + ausgefüllte citations_template.json und:

- entfernt die Marker ⟦CITATION?⟧ / ⟦CITATION?/SENTIMENT⟧ / ⟦SENTIMENT⟧
- setzt hochgestellte Fussnotenzahlen an die entsprechenden Sätze
- hängt am Dokumentende einen Abschnitt "Fussnoten" mit den Texten an

Input:
    INPUT_DOCX            – dein von hist_docx_evidence_marker.py erzeugtes DOCX
    CITATIONS_FILLED_JSON – citations_template.json, in dem "footnote" gefüllt ist
Output:
    OUTPUT_DOCX           – neues DOCX mit nummerierten Verweisen + Fussnotenliste
"""

from pathlib import Path
import json
import re
from typing import List, Dict, Tuple

from docx import Document
from docx.text.run import Run

# Konfiguration
INPUT_DOCX = Path("output_marked.docx")
CITATIONS_FILLED_JSON = Path("citations_template.json")  # nach dem Füllen
OUTPUT_DOCX = Path("output_with_footnotes.docx")

# Marker, die dein erstes Script einfügt
MARKER_PREFIXES = (
    "⟦CITATION?/SENTIMENT⟧",
    "⟦CITATION?⟧",
    "⟦SENTIMENT⟧",
)


def split_into_sentences(text: str) -> List[str]:
    """
    Sehr einfache Satzsegmentierung (wie im ersten Script):
    Splittet an . ! ? + Leerzeichen.
    """
    text = text.strip()
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p.strip() for p in parts if p.strip()]


def load_citations() -> Dict[Tuple[int, int], str]:
    """
    Lädt citations_template.json und gibt Mapping
    (paragraph_index, sentence_index) -> footnote_text zurück.
    """
    if not CITATIONS_FILLED_JSON.exists():
        raise FileNotFoundError(f"{CITATIONS_FILLED_JSON} nicht gefunden")

    with CITATIONS_FILLED_JSON.open("r", encoding="utf-8") as f:
        items = json.load(f)

    mapping: Dict[Tuple[int, int], str] = {}
    for entry in items:
        foot = (entry.get("footnote") or "").strip()
        if not foot:
            # Eintrag ohne Fussnote wird ignoriert
            continue
        key = (int(entry["paragraph_index"]), int(entry["sentence_index"]))
        mapping[key] = foot
    return mapping


def strip_leading_marker(sentence: str) -> str:
    """
    Entfernt einen der Marker am Satzanfang, falls vorhanden.
    """
    s = sentence.lstrip()
    for marker in MARKER_PREFIXES:
        if s.startswith(marker):
            # Marker + evtl. folgendes Leerzeichen entfernen
            return s[len(marker) :].lstrip()
    return sentence


def process_docx():
    if not INPUT_DOCX.exists():
        raise FileNotFoundError(f"{INPUT_DOCX} nicht gefunden")

    citations = load_citations()
    doc = Document(str(INPUT_DOCX))

    # Liste aller Fussnoten in Reihenfolge des Auftretens
    footnotes_ordered: List[Tuple[int, int, str]] = []  # (paragraph_index, sentence_index, text)

    next_number = 1

    for p_idx, para in enumerate(doc.paragraphs):
        original = para.text
        if not original.strip():
            continue

        sentences = split_into_sentences(original)
        if not sentences:
            continue

        new_runs: List[Run] = []

        for s_idx, sent in enumerate(sentences, start=1):
            # Hat dieser Satz eine Fussnote?
            key = (p_idx, s_idx)
            needs_note = key in citations

            clean_sentence = strip_leading_marker(sent)

            # Satz als Run einfügen
            r_sentence = para._element._new_r()
            run_obj = Run(r_sentence, para)
            run_obj.text = clean_sentence
            new_runs.append(run_obj)

            if needs_note:
                # Nummer zuweisen
                number = next_number
                next_number += 1
                footnotes_ordered.append((p_idx, s_idx, citations[key]))

                # Hochgestellte Nummer direkt nach dem Satz
                r_sup = para._element._new_r()
                run_sup = Run(r_sup, para)
                run_sup.text = str(number)
                run_sup.font.superscript = True
                new_runs.append(run_sup)

            # Zwischen Sätzen wieder ein Leerzeichen
            if s_idx != len(sentences):
                r_space = para._element._new_r()
                run_space = Run(r_space, para)
                run_space.text = " "
                new_runs.append(run_space)

        # Alte Runs löschen und neue einsetzen
        # (einfachste Methode: Paragraph leeren, dann Runs anhängen)
        para.clear()
        for r in new_runs:
            para._p.append(r._r)

    # Abschnitt mit Fussnoten anhängen
    if footnotes_ordered:
        doc.add_page_break()
        heading = doc.add_paragraph()
        heading.style = "Heading 1"
        heading.add_run("Fussnoten")

        for idx, (_, _, text) in enumerate(footnotes_ordered, start=1):
            p = doc.add_paragraph()
            num_run = p.add_run(f"{idx} ")
            num_run.bold = True
            p.add_run(text)

    doc.save(str(OUTPUT_DOCX))
    print(f"[OK] {OUTPUT_DOCX} geschrieben mit {len(footnotes_ordered)} Fussnoten.")


if __name__ == "__main__":
    process_docx()